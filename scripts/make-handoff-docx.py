"""Regenerate the Drive copy of the mobile handoff from docs/HANDOFF.md.

    python scripts/make-handoff-docx.py "G:/Shared drives/07-OHRR App/OHRR Mobile Build Handoff.docx"

Plain Markdown → Word: headings, bullet / numbered lists, tables, `code`,
**bold**, links kept as text. Needs python-docx (pip install python-docx).
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

SRC = Path(__file__).resolve().parent.parent / "docs" / "HANDOFF.md"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else SRC.with_suffix(".docx")

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            r = par.add_run(f"{m.group(1)} ({m.group(2)})" if m and m.group(1) != m.group(2) else (m.group(2) if m else part))
        else:
            par.add_run(part)


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:-]+\|", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            rows = [r for k, r in enumerate(rows) if k != 1]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row[: len(rows[0])]):
                    p = table.cell(r, c).paragraphs[0]
                    add_inline(p, cell)
                    if r == 0:
                        for run in p.runs:
                            run.bold = True
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
        elif line.startswith("> "):
            p = doc.add_paragraph()
            add_inline(p, line[2:])
            for run in p.runs:
                run.italic = True
        elif re.match(r"^\s*[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\s*\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\s*\d+\.\s+", "", line))
        elif line.strip() == "":
            pass
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
        i += 1
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
